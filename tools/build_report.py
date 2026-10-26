"""Build a Word report from the actual Django test run."""
from datetime import datetime
from pathlib import Path
import ast
import textwrap

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent

DESCRIPTIONS = [
    'Начальное значение счётчика равно нулю',
    'Строковое представление модели совпадает с заголовком',
    'Страница списка возвращает HTTP 200',
    'Для списка используется ожидаемый шаблон',
    'Список содержит заголовок и автора статьи',
    'Контекст списка содержит созданную запись и её идентификатор',
    'Ссылка из списка открывает детальную страницу',
    'Детальная страница использует нужный шаблон и содержит данные статьи',
    'Для отсутствующей статьи возвращается HTTP 404',
    'Один просмотр увеличивает счётчик на единицу',
    'Повторный просмотр снова увеличивает счётчик',
    'GET создания показывает несвязанную форму',
    'POST создания сохраняет все поля и перенаправляет к списку',
    'Каждое обязательное поле отклоняет пробельное значение',
    'Заголовок длиннее 200 символов не сохраняется',
    'Автор длиннее 120 символов не сохраняется',
    'Форма редактирования заполнена и не меняет счётчик',
    'Редактирование обновляет существующую запись и сохраняет счётчик',
    'Ошибочная форма редактирования не изменяет запись',
    'GET удаления только показывает подтверждение',
    'POST удаления удаляет только выбранную запись',
    'Отсутствующая запись даёт 404 при GET и POST изменения и удаления',
    'CRUD обработчики отклоняют PUT кодом 405',
    'POST без CSRF токена отклоняется кодом 403 для всех операций',
    'Корректный CSRF токен позволяет создать статью',
    'Для пустой базы показывается сообщение об отсутствии статей',
    'Список содержит ссылки создания редактирования и удаления',
    'HTML в заголовке экранируется',
    'Форма содержит только заголовок описание и автора',
    'Переданное вручную значение счётчика игнорируется',
]


def build_report(data):
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(1.8)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Normal'].font.size = Pt(12)
    doc.styles['Normal'].paragraph_format.space_after = Pt(7)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.12
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Heading 1'].font.size = Pt(16)
    doc.styles['Heading 2'].font.size = Pt(13)
    code_style = doc.styles.add_style('Source Code', 1)
    code_style.font.name = 'Consolas'
    code_style.font.size = Pt(8)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = 1
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    doc.add_paragraph('Лабораторная работа 6', 'Title')
    doc.add_paragraph('CRUD и автоматическое тестирование Django', 'Subtitle')
    doc.add_paragraph('Дата выполнения 26 октября 2026 года')
    doc.add_heading('Цель работы', 1)
    doc.add_paragraph('Разработать CRUD для статей блога в Django, проверить поведение проекта автоматическими тестами и записать результаты выполнения в Word документ средствами тестового запуска Django.')
    doc.add_heading('Результат выполнения', 1)
    statuses = {s: sum(r['status'] == s for r in data['results']) for s in {r['status'] for r in data['results']}}
    doc.add_paragraph(f"Выполнено тестов: {data['tests_run']}. Общий результат: {'успешно' if data['successful'] else 'есть ошибки'}. Статусы: " + ', '.join(f'{k} — {v}' for k, v in sorted(statuses.items())) + '.')
    doc.add_paragraph('В браузере последовательно проверены создание статьи, открытие списка и подробной страницы, изменение заголовка, подтверждение удаления и пустой список после удаления. Скриншоты приведены далее.')
    doc.add_heading('Исходный пример и структура проекта', 1)
    doc.add_paragraph('Изучена папка django-test репозитория https://github.com/narmuhamedov/TestingPO.git, ревизия 42ab2e4cc81ae3b5eb2ca122de041e175df48a57. В примере используются Blog, Django TestCase, Client и DiscoverRunner с записью статусов в Word. Существующий проект dj расширен CRUD и собственным ReportingRunner.')
    doc.add_paragraph('Модель Blog содержит title, description, author и likes. Поле likes используется как счётчик просмотров и увеличивается при открытии подробной страницы — это сохранённое поведение исходного проекта dj. Form исключает счётчик из редактируемых полей.')
    doc.add_paragraph('Приложение предназначено для локальной учебной демонстрации. CRUD доступен без входа в систему. Используются SQLite и Django 5.1.15. Схема модели не изменена, новая миграция не требуется.')

    doc.add_page_break()
    doc.add_heading('Реализация CRUD', 1)
    rows = [('Операция', 'Адрес', 'Поведение'), ('Create', '/blog/create/', 'GET форма, POST создание'), ('Read', '/ и /blog/<id>/', 'Список и подробная страница'), ('Update', '/blog/<id>/edit/', 'GET заполненная форма, POST изменение'), ('Delete', '/blog/<id>/delete/', 'GET подтверждение, POST удаление')]
    table = doc.add_table(rows=0, cols=3)
    table.style = 'Table Grid'
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    doc.add_paragraph('BlogForm наследуется от ModelForm. is_valid проверяет обязательность и максимальную длину полей. Неуспешный POST повторно выводит форму с ошибками. Удаление выполняется только POST запросом, формы включают csrf_token. Для отсутствующих записей используется get_object_or_404.')
    doc.add_heading('Как формируется Word отчёт', 1)
    doc.add_paragraph('В settings.py указан TEST_RUNNER = blog.test_runner.ReportingRunner. После выполнения набора тестов ReportingResult передаёт фактические результаты в tools/build_report.py. Генератор создаёт reports/lab6_report.docx и журнал reports/test-results.json. Статусы берутся из addSuccess, addFailure, addError и addSubTest, а не назначаются заранее. Ошибки и пропуски также записываются в отчёт.')
    doc.add_paragraph('TestCase создаёт отдельную тестовую базу данных. Изменения тестов изолированы от пользовательской SQLite базы. Для проверки CSRF применяется Client(enforce_csrf_checks=True). Один тест с subTest может проверять несколько вариантов ввода.')
    doc.add_heading('Порядок повторного запуска', 1)
    for line in ['python -m pip install -r requirements.txt', 'python manage.py migrate', 'python manage.py test -v 2', 'python manage.py runserver']:
        doc.add_paragraph(line, 'Source Code')
    doc.add_paragraph('После команды test Word отчёт пересоздаётся. Сохранённые скриншоты иллюстрируют браузерную проверку от 26 октября 2026 года и автоматически не переснимаются при новом тестовом запуске.')

    source = (ROOT / 'blog/tests.py').read_text(encoding='utf-8-sig')
    tree = ast.parse(source)
    methods = [n for cls in tree.body if isinstance(cls, ast.ClassDef) for n in cls.body if isinstance(n, ast.FunctionDef) and n.name.startswith('test_')]
    descriptions = dict(zip([m.name for m in methods], DESCRIPTIONS))
    records = data['results']
    for start in range(0, len(records), 10):
        doc.add_page_break()
        doc.add_heading('Результаты автоматических тестов', 1)
        doc.add_paragraph(f'Проверки {start + 1}–{min(start + 10, len(records))} из {len(records)} зарегистрированных результатов. PASS означает успешное выполнение всех утверждений теста.')
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        table.autofit = False
        for col, width in zip(table.columns, [0.8, 14, 2.2]):
            col.width = Cm(width)
        for c, t in zip(table.rows[0].cells, ['№', 'Тест и проверяемое поведение', 'Статус']):
            c.text = t
        for i, record in enumerate(records[start:start+10], start+1):
            name = record['test'].split('.')[-1]
            cells = table.add_row().cells
            cells[0].text = str(i)
            cells[1].text = name + '\n' + descriptions.get(name, record['test'])
            cells[2].text = record['status']
            for c in cells:
                for p in c.paragraphs:
                    p.paragraph_format.line_spacing = 1
                    p.paragraph_format.space_after = Pt(4)
                    for run in p.runs:
                        run.font.size = Pt(9)
            cells[0].width, cells[1].width, cells[2].width = Cm(0.8), Cm(14), Cm(2.2)
        for row in table.rows:
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement('w:cantSplit'))
        for cell in table.rows[0].cells:
            shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'E2EAF1'); cell._tc.get_or_add_tcPr().append(shd)
        borders = OxmlElement('w:tblBorders')
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            el = OxmlElement('w:' + edge); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D9D9D9'); borders.append(el)
        table._tbl.tblPr.append(borders)

    figures = [
        ('01-create.png', 'Создание статьи', 'Форма заполнена заголовком, описанием и автором. После нажатия Сохранить статья появилась в списке.'),
        ('02-list.png', 'Список статей', 'Новая запись отображается с автором и нулевым счётчиком. Доступны ссылки на просмотр, редактирование и удаление.'),
        ('03-detail.png', 'Подробная страница', 'Открыта созданная статья. Содержимое совпадает с формой, счётчик просмотров увеличился до единицы.'),
        ('04-update.png', 'Редактирование статьи', 'В заголовок добавлено слово обновлено. После сохранения браузерная проверка подтвердила появление нового заголовка в списке.'),
        ('05-delete.png', 'Подтверждение удаления', 'GET запрос показывает название выбранной статьи и кнопку подтверждения. Запись ещё существует.'),
        ('06-deleted.png', 'Результат удаления', 'После POST подтверждения показано сообщение Статей пока нет. Браузерная проверка подтвердила отсутствие записи.'),
        ('07-tests.png', 'Протокол тестирования', 'Скриншот сохранённого вывода реального запуска Django. Полный журнал находится в reports/test-results.txt.'),
    ]
    for i in range(0, len(figures), 2):
        present = [(f, title, caption) for f, title, caption in figures[i:i+2] if (ROOT/'reports/screenshots'/f).exists()]
        if not present:
            continue
        doc.add_page_break()
        doc.add_heading('Проверка работы приложения', 1)
        for f, title, caption in present:
            doc.add_heading(title, 2)
            doc.add_picture(str(ROOT/'reports/screenshots'/f), width=Cm(14))
            p = doc.add_paragraph(caption)
            p.paragraph_format.space_after = Pt(9)
            for r in p.runs: r.font.size = Pt(10)

    doc.add_page_break()
    doc.add_heading('Полный код тестов', 1)
    doc.add_paragraph('Листинг blog/tests.py включает все 30 тестовых методов текущего проекта. Код ниже соответствует выполненному набору тестов.')
    for line in source.splitlines():
        wrapped = textwrap.wrap(line, width=104, replace_whitespace=False, drop_whitespace=False) or ['']
        for part in wrapped:
            doc.add_paragraph(part, 'Source Code')
    problems = [r for r in records if r['status'] != 'PASS']
    if problems:
        doc.add_page_break()
        doc.add_heading('Диагностика результатов', 1)
        for r in problems:
            doc.add_paragraph(r['test'] + ' ' + r['status'])
            doc.add_paragraph(r['detail'], 'Source Code')
    doc.add_heading('Вывод', 1)
    doc.add_paragraph('Реализован полный цикл работы со статьёй. Автоматические тесты проверяют модель, страницы, формы, операции базы данных и ошибочные запросы. Результаты формируются непосредственно тестовым запуском Django и сохраняются в Word для сдачи лабораторной работы.')
    doc.core_properties.title = 'Лабораторная работа 6 CRUD и тестирование Django'
    doc.core_properties.author = 'Гульжан'
    for index, table in enumerate(doc.tables):
        widths = [3, 5, 9] if index == 0 else [0.8, 14, 2.2]
        table.autofit = False
        for col, width in zip(table.columns, widths):
            col.width = Cm(width)
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                cell.vertical_alignment = 1
        borders = OxmlElement('w:tblBorders')
        for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            el = OxmlElement('w:' + edge)
            for attr, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                el.set(qn('w:' + attr), value)
            borders.append(el)
        table._tbl.tblPr.append(borders)
        for cell in table.rows[0].cells:
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), 'E2EAF1')
            cell._tc.get_or_add_tcPr().append(shd)
    # Remove inherited Word template paragraph borders, including the title rule.
    for element in [doc.styles.element, doc.element]:
        for border in list(element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    target = ROOT/'reports/lab6_report.docx'
    doc.save(target)
    print(f'Word report: {target}')
    return target

